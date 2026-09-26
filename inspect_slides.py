import pptx

prs = pptx.Presentation('presentation_downloaded.pptx')
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    title = ""
    texts = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
    if texts:
        title = texts[0]
    print(f"Slide {idx+1}: {title} (Total text elements: {len(texts)})")
